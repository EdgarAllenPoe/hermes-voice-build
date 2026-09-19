#!/usr/bin/env python3
"""Build printable DOCX and matching Markdown from docs/hardware-guide.json."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
ROOT=Path(__file__).resolve().parents[1]
def table(doc,block):
    rows=block['rows'];widths=block.get('widths',[6.7/len(rows[0])]*len(rows[0]))
    t=doc.add_table(rows=0,cols=len(rows[0]));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    for index,row in enumerate(rows):
        cells=t.add_row().cells
        for col,text in enumerate(row):
            c=cells[col];c.width=Inches(widths[col]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.space_before=Pt(1)
            p.paragraph_format.line_spacing=1.03
            run=p.add_run(text);run.font.size=Pt(9.4);run.bold=index==0
            props=c._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
            for side in ('top','left','bottom','right'):
                edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
            props.append(borders);margins=OxmlElement('w:tcMar')
            for side,value in [('top','65'),('bottom','65'),('left','95'),('right','95')]:
                e=OxmlElement('w:'+side);e.set(qn('w:w'),value);e.set(qn('w:type'),'dxa');margins.append(e)
            props.append(margins)
            if index==0:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8E8E8');props.append(shade)
        rowprops=t.rows[-1]._tr.get_or_add_trPr();rowprops.append(OxmlElement('w:cantSplit'))
        if index==0:rowprops.append(OxmlElement('w:tblHeader'))
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
def main():
    pages=json.loads((ROOT/'docs/hardware-guide.json').read_text(encoding='utf-8'))
    doc=Document();section=doc.sections[0]
    for node in doc.styles.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    section.page_width=Inches(8.5);section.page_height=Inches(11)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.8);section.footer_distance=Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2']:
        s=doc.styles[name];s.font.name='Calibri';s.font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.paragraph_format.space_after=Pt(6);normal.paragraph_format.line_spacing=1.08
    doc.styles['Title'].font.size=Pt(25);doc.styles['Heading 1'].font.size=Pt(18)
    doc.styles['Heading 1'].paragraph_format.space_after=Pt(10)
    doc.styles['Heading 2'].font.size=Pt(12);doc.styles['Heading 2'].paragraph_format.space_before=Pt(8)
    footer=section.footer.paragraphs[0]
    footer.add_run('Hermes Voice firmware 0.3.3   |   Hardware acceptance pending   |   Page ').font.size=Pt(8)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    md=[]
    for index,page in enumerate(pages):
        if index:doc.add_page_break()
        doc.add_paragraph(page['title'],'Title' if index==0 else 'Heading 1')
        md.append('# '+page['title']+'\n')
        for block in page['blocks']:
            kind=block['type']
            if kind=='text':
                p=doc.add_paragraph(block['text'])
                if block.get('bold'):
                    for run in p.runs:run.bold=True
                md.append(block['text']+'\n')
            elif kind=='heading':
                doc.add_paragraph(block['text'],'Heading 2');md.append('## '+block['text']+'\n')
            elif kind=='steps':
                for i,text in enumerate(block['items'],1):
                    p=doc.add_paragraph(str(i)+'. '+text)
                    p.paragraph_format.left_indent=Inches(.17);p.paragraph_format.first_line_indent=Inches(-.17)
                    md.append(str(i)+'. '+text)
                md.append('')
            elif kind=='code':
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(7);p.paragraph_format.line_spacing=1
                run=p.add_run(block['text']);run.font.name='Consolas';run.font.size=Pt(8.1)
                md+=[chr(96)*3+block.get('language','text'),block['text'],chr(96)*3+'\n']
            elif kind=='table':
                table(doc,block);rows=block['rows']
                md+=['| '+' | '.join(rows[0])+' |','| '+' | '.join(['---']*len(rows[0]))+' |']
                md+=['| '+' | '.join(row)+' |' for row in rows[1:]];md.append('')
            elif kind=='image':
                p=doc.add_paragraph();p.add_run().add_picture(str(ROOT/'docs'/block['file']),width=Inches(block.get('width',6.7)))
                drawing=p._p.xpath('.//wp:docPr')[0];drawing.set('descr',block['alt'])
                md.append('!['+block['alt']+']('+block['file']+')\n')
    doc.core_properties.title='Hermes Voice hardware workshop guide'
    doc.core_properties.subject='Wiring programming and commissioning for XIAO nRF54LM20A Sense'
    doc.core_properties.author='Hermes Voice Project';doc.core_properties.keywords='Hermes Voice firmware 0.3.3 hardware commissioning'
    out=ROOT/'docs/Hermes-Voice-Hardware-Guide-0.3.3.docx';doc.save(out)
    (ROOT/'docs/HARDWARE-GUIDE.md').write_text('\n'.join(md),encoding='utf-8',newline='\n')
    print('Created '+str(out)+' from '+str(len(pages))+' planned workshop pages')
if __name__=='__main__':main()
